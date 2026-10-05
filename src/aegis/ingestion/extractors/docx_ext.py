from pathlib import Path
from typing import List
from docx import Document as DocxDocument
from aegis.models.domain import Document, ExtractionMethod, ExtractionReliability, ExtractionStatus, LocationType
from .base import DocumentExtractor, ExtractedUnit

class DocxExtractor(DocumentExtractor):
    def extract(self, document: Document, file_path: Path) -> List[ExtractedUnit]:
        units = []
        try:
            doc = DocxDocument(file_path)
            # Paragraphs
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    clean_text = " ".join(text.split())
                    units.append(ExtractedUnit(
                        raw_text=clean_text,
                        location_type=LocationType.SECTION,
                        location_value=f"Paragraph {i}",
                        reliability=ExtractionReliability.HIGH,
                        status=ExtractionStatus.SUCCESS,
                        method=ExtractionMethod.DETERMINISTIC
                    ))
            # Tables
            for t_idx, table in enumerate(doc.tables):
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        text = cell.text.strip()
                        if text:
                            clean_text = " ".join(text.split())
                            units.append(ExtractedUnit(
                                raw_text=clean_text,
                                location_type=LocationType.CELL,
                                location_value=f"Table {t_idx} Row {r_idx} Col {c_idx}",
                                reliability=ExtractionReliability.HIGH,
                                status=ExtractionStatus.SUCCESS,
                                method=ExtractionMethod.DETERMINISTIC
                            ))
        except Exception as e:
            units.append(ExtractedUnit(
                raw_text="",
                location_type=LocationType.SECTION,
                location_value="Unknown",
                reliability=ExtractionReliability.LOW,
                status=ExtractionStatus.FAILED,
                method=ExtractionMethod.DETERMINISTIC
            ))
        return units
