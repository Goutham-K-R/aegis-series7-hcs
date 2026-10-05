from pathlib import Path
from typing import List
import openpyxl
from aegis.models.domain import Document, ExtractionMethod, ExtractionReliability, ExtractionStatus, LocationType
from .base import DocumentExtractor, ExtractedUnit

class XlsxExtractor(DocumentExtractor):
    def extract(self, document: Document, file_path: Path) -> List[ExtractedUnit]:
        units = []
        try:
            wb = openpyxl.load_workbook(filename=file_path, data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                headers = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if row_idx == 1:
                        headers = [str(c).strip() if c else f"Col_{i}" for i, c in enumerate(row)]
                        continue
                        
                    has_data = any(c is not None and str(c).strip() for c in row)
                    if not has_data:
                        continue
                        
                    row_data = {}
                    for col_idx, cell_value in enumerate(row):
                        if cell_value is not None and str(cell_value).strip():
                            col_name = headers[col_idx] if col_idx < len(headers) else f"Col_{col_idx}"
                            row_data[col_name] = str(cell_value).strip()
                            
                    if row_data:
                        import json
                        raw = json.dumps(row_data)
                        loc = f"Sheet: {sheet_name}, Row: {row_idx}"
                        units.append(ExtractedUnit(
                            raw_text=raw,
                            location_type=LocationType.SHEET,
                            location_value=loc,
                            reliability=ExtractionReliability.HIGH,
                            status=ExtractionStatus.SUCCESS,
                            method=ExtractionMethod.DETERMINISTIC
                        ))
            wb.close()
        except Exception as e:
            units.append(ExtractedUnit(
                raw_text="",
                location_type=LocationType.SHEET,
                location_value="Unknown",
                reliability=ExtractionReliability.LOW,
                status=ExtractionStatus.FAILED,
                method=ExtractionMethod.DETERMINISTIC
            ))
        return units
