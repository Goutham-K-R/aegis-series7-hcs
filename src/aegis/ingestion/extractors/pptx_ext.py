from pathlib import Path
from typing import List
from pptx import Presentation
from aegis.models.domain import Document, ExtractionMethod, ExtractionReliability, ExtractionStatus, LocationType
from .base import DocumentExtractor, ExtractedUnit

class PptxExtractor(DocumentExtractor):
    def extract(self, document: Document, file_path: Path) -> List[ExtractedUnit]:
        units = []
        try:
            prs = Presentation(file_path)
            for s_idx, slide in enumerate(prs.slides):
                for shp_idx, shape in enumerate(slide.shapes):
                    if hasattr(shape, "text") and shape.text.strip():
                        clean_text = " ".join(shape.text.split())
                        units.append(ExtractedUnit(
                            raw_text=clean_text,
                            location_type=LocationType.SLIDE,
                            location_value=f"Slide {s_idx + 1} Shape {shp_idx}",
                            reliability=ExtractionReliability.HIGH,
                            status=ExtractionStatus.SUCCESS,
                            method=ExtractionMethod.DETERMINISTIC
                        ))
                    if shape.has_table:
                        for r_idx, row in enumerate(shape.table.rows):
                            for c_idx, cell in enumerate(row.cells):
                                text = cell.text_frame.text.strip()
                                if text:
                                    clean_text = " ".join(text.split())
                                    units.append(ExtractedUnit(
                                        raw_text=clean_text,
                                        location_type=LocationType.CELL,
                                        location_value=f"Slide {s_idx + 1} Table {shp_idx} Row {r_idx} Col {c_idx}",
                                        reliability=ExtractionReliability.HIGH,
                                        status=ExtractionStatus.SUCCESS,
                                        method=ExtractionMethod.DETERMINISTIC
                                    ))
        except Exception as e:
            units.append(ExtractedUnit(
                raw_text="",
                location_type=LocationType.SLIDE,
                location_value="Unknown",
                reliability=ExtractionReliability.LOW,
                status=ExtractionStatus.FAILED,
                method=ExtractionMethod.DETERMINISTIC
            ))
        return units
