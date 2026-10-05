import pytest
import openpyxl
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pathlib import Path
from aegis.models.domain import Document, Format
from aegis.ingestion.extractors.json_ext import JsonExtractor
from aegis.ingestion.extractors.xlsx_ext import XlsxExtractor
from aegis.ingestion.extractors.html_ext import HtmlExtractor
from aegis.ingestion.extractors.docx_ext import DocxExtractor
from aegis.ingestion.extractors.image_ext import ImageExtractor

def test_json_extractor(tmp_path):
    p = tmp_path / "test.json"
    p.write_bytes(b'{"system": {"version": "3.2.1", "pressure": "180 bar"}}')
    
    doc = Document(filename="test.json", relative_path="test.json", format=Format.JSON, hash="123", file_size=100)
    ext = JsonExtractor()
    units = ext.extract(doc, p)
    
    assert len(units) == 2
    assert any(u.raw_text == "3.2.1" and u.location_value == "$.system.version" for u in units)
    assert any(u.raw_text == "180 bar" and u.location_value == "$.system.pressure" for u in units)

def test_conflict_safety_in_json_extractor(tmp_path):
    p = tmp_path / "test.json"
    # Document contains "180 bar". We must extract "180 bar" exactly.
    # We must NOT rewrite it to "200 bar" based on outside knowledge.
    p.write_bytes(b'{"pressure": "180 bar"}')
    doc = Document(filename="test.json", relative_path="test.json", format=Format.JSON, hash="123", file_size=100)
    units = JsonExtractor().extract(doc, p)
    assert units[0].raw_text == "180 bar"

def test_html_extractor(tmp_path):
    p = tmp_path / "test.html"
    p.write_bytes(b'<html><body><h1>Title</h1><p>Some   text\n here</p></body></html>')
    
    doc = Document(filename="test.html", relative_path="test.html", format=Format.HTML, hash="123", file_size=100)
    ext = HtmlExtractor()
    units = ext.extract(doc, p)
    
    assert len(units) == 2
    assert units[0].raw_text == "Title"
    assert units[1].raw_text == "Some text here" # Tests whitespace normalization

def test_image_extractor(tmp_path):
    p = tmp_path / "test.png"
    p.write_bytes(b"fake png")
    
    doc = Document(filename="test.png", relative_path="test.png", format=Format.PNG, hash="123", file_size=100)
    ext = ImageExtractor()
    units = ext.extract(doc, p)
    
    assert len(units) == 1
    assert units[0].status.value == "REQUIRES_OCR"

def test_xlsx_extractor(tmp_path):
    p = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conf"
    ws['B7'] = "3.2.1"
    wb.save(p)
    
    doc = Document(filename="test.xlsx", relative_path="test.xlsx", format=Format.XLSX, hash="123", file_size=100)
    units = XlsxExtractor().extract(doc, p)
    
    assert len(units) == 1
    assert '3.2.1' in units[0].raw_text
    assert units[0].location_value == "Sheet: Conf, Row: 7"

def test_docx_extractor(tmp_path):
    p = tmp_path / "test.docx"
    d = DocxDocument()
    d.add_paragraph("Paragraph 0 text")
    d.save(p)
    
    doc = Document(filename="test.docx", relative_path="test.docx", format=Format.DOCX, hash="123", file_size=100)
    units = DocxExtractor().extract(doc, p)
    
    assert len(units) == 1
    assert units[0].raw_text == "Paragraph 0 text"
